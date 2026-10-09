from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from matplotlib.figure import Figure
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

from tests.helpers import load_script, study_tables

I = load_script("04_mrmc_inference_teaching.py")


class InferenceTests(unittest.TestCase):
    def setUp(self):
        self.data = study_tables()[2]
        self.readers, self.cases = I.validate_input(self.data)
        self.arrays = I.build_metric_arrays(self.data, self.readers, self.cases)

    def test_hand_calculated_reader_average(self):
        metrics = I.calculate_reader_averaged_metrics(self.arrays, np.ones(2), np.ones(3))
        expected = {"lesion_sensitivity_without": 2/3, "lesion_sensitivity_with": 5/6,
                    "fp_per_scan_without": .5, "fp_per_scan_with": 2/3,
                    "reading_time_min_without": 5, "reading_time_min_with": 3.5}
        for metric, value in expected.items():
            self.assertAlmostEqual(metrics[metric], value)

    def test_weighted_repeated_reader_and_cases(self):
        metrics = I.calculate_reader_averaged_metrics(self.arrays, np.array([0., 2.]), np.array([1., 0., 2.]))
        # Only R2, cases C1/C3/C3: without=9+3+3, with=7+1+1.
        self.assertAlmostEqual(metrics["lesion_sensitivity_without"], 1)
        self.assertAlmostEqual(metrics["lesion_sensitivity_with"], 1)
        self.assertAlmostEqual(metrics["fp_per_scan_without"], 1)
        self.assertAlmostEqual(metrics["fp_per_scan_with"], 1/3)
        self.assertAlmostEqual(metrics["reading_time_min_without"], 5)
        self.assertAlmostEqual(metrics["reading_time_min_with"], 3)

    def test_unsorted_input_builds_same_arrays(self):
        shuffled = self.data.sample(frac=1, random_state=7)
        arrays = I.build_metric_arrays(shuffled, *I.validate_input(shuffled))
        for key in self.arrays:
            np.testing.assert_array_equal(arrays[key], self.arrays[key])

    def test_missing_column_missing_value_and_empty_table(self):
        missing_value = self.data.copy()
        missing_value.loc[0, "reader_id"] = None
        for table in (self.data.drop(columns="fp_count"), missing_value, self.data.iloc[:0]):
            with self.subTest(rows=len(table), columns=len(table.columns)), self.assertRaises(ValueError):
                I.validate_input(table)

    def test_duplicate_incomplete_pair_and_invalid_conditions(self):
        invalid_condition = self.data.copy()
        invalid_condition.loc[0, "condition"] = "other"
        for table in (pd.concat([self.data, self.data.iloc[[0]]]), self.data.iloc[1:], invalid_condition):
            with self.subTest(rows=len(table)), self.assertRaises(ValueError):
                I.validate_input(table)

    def test_period_and_tp_fn_must_be_valid(self):
        for column, value in [("period", 3), ("period", 2), ("fn_count", 100)]:
            table = self.data.copy()
            table.loc[0, column] = value
            with self.subTest(column=column, value=value), self.assertRaises(ValueError):
                I.validate_input(table)

    def test_counts_and_times_must_be_finite_numeric_and_valid(self):
        for column, value in [("fp_count", -1), ("tp_count", .5), ("fp_count", .5), ("fp_count", np.inf),
                              ("reading_time_min", 0), ("reading_time_min", -1),
                              ("reading_time_min", np.inf), ("reading_time_min", "bad")]:
            table = self.data.copy()
            table[column] = table[column].astype(object)
            table.loc[0, column] = value
            with self.subTest(column=column, value=value), self.assertRaises(ValueError):
                I.validate_input(table)

    def test_case_reference_must_be_same_across_readers_and_conditions(self):
        table = self.data.copy()
        table.loc[0, "reference_lesion_count"] = 3
        table.loc[0, "fn_count"] = 2  # TP+FN still agrees locally, but C1 differs.
        with self.assertRaisesRegex(ValueError, "参考"):
            I.validate_input(table)

    def test_weights_and_zero_reference_resample_are_rejected(self):
        for reader_weights, case_weights in [(np.zeros(2), np.ones(3)), (np.ones(2), np.zeros(3)),
                                             (np.ones(2), np.array([0., 0., 3.]))]:
            with self.subTest(weights=case_weights), self.assertRaises(ValueError):
                I.calculate_reader_averaged_metrics(self.arrays, reader_weights, case_weights)

    def test_negative_nonfinite_and_wrong_shape_weights_are_rejected(self):
        for reader_weights, case_weights in [(np.array([-1., 3.]), np.ones(3)), (np.array([np.nan, 1]), np.ones(3)),
                                             (np.ones(2), np.array([np.inf, 1, 1])), (np.ones(1), np.ones(3))]:
            with self.subTest(reader_weights=reader_weights), self.assertRaises(ValueError):
                I.calculate_reader_averaged_metrics(self.arrays, reader_weights, case_weights)

    def test_fixed_seed_reproduces_bootstrap(self):
        # All-positive tiny array avoids undefined all-negative-case resamples.
        arrays = {key: value[:, :2] for key, value in self.arrays.items()}
        with patch.object(I, "BOOTSTRAP_ITERATIONS", 30):
            first_point, first = I.run_two_way_cluster_bootstrap(arrays, 2, 2)
            second_point, second = I.run_two_way_cluster_bootstrap(arrays, 2, 2)
        self.assertEqual(first_point, second_point)
        pd.testing.assert_frame_equal(first, second)
        self.assertEqual(first.iteration.tolist(), list(range(1, 31)))

    def test_controlled_resampling_uses_same_reader_case_weights_for_both_conditions(self):
        class FixedRng:
            def __init__(self):
                self.samples = iter([np.array([1, 1]), np.array([2, 0, 2])])

            def integers(self, low, high, size):
                sample = next(self.samples)
                assert len(sample) == size and high in (2, 3)
                return sample

        with patch.object(I, "BOOTSTRAP_ITERATIONS", 1), patch.object(I.np.random, "default_rng", return_value=FixedRng()):
            _, distribution = I.run_two_way_cluster_bootstrap(self.arrays, 2, 3)
        self.assertAlmostEqual(distribution.iloc[0].lesion_sensitivity_difference, 0)
        self.assertAlmostEqual(distribution.iloc[0].fp_per_scan_difference, -2/3)
        self.assertAlmostEqual(distribution.iloc[0].reading_time_difference_min, -2)

    def test_constant_paired_effect_is_preserved_in_all_resamples(self):
        arrays = {key: value[:, :2].copy() for key, value in self.arrays.items()}
        arrays["reading_time_min_with"] = arrays["reading_time_min_without"] - 1
        arrays["fp_count_with"] = arrays["fp_count_without"] + 2
        arrays["tp_count_with"] = arrays["tp_count_without"].copy()
        with patch.object(I, "BOOTSTRAP_ITERATIONS", 30):
            _, distribution = I.run_two_way_cluster_bootstrap(arrays, 2, 2)
        np.testing.assert_allclose(distribution.reading_time_difference_min, -1)
        np.testing.assert_allclose(distribution.fp_per_scan_difference, 2)
        np.testing.assert_allclose(distribution.lesion_sensitivity_difference, 0)

    def test_summary_se_percentile_interval_and_pvalue(self):
        point = {"lesion_sensitivity_without": .5, "lesion_sensitivity_with": .75,
                 "fp_per_scan_without": .5, "fp_per_scan_with": .75,
                 "reading_time_min_without": 2, "reading_time_min_with": 2.25}
        distribution = pd.DataFrame({"lesion_sensitivity_difference": [0, 1, 2],
                                     "fp_per_scan_difference": [0, 1, 2], "reading_time_difference_min": [0, 1, 2]})
        summary = I.summarize_inference(point, distribution)
        np.testing.assert_allclose(summary.difference, .25)
        np.testing.assert_allclose(summary.bootstrap_se, 1)
        np.testing.assert_allclose(summary.ci_lower, .05)
        np.testing.assert_allclose(summary.ci_upper, 1.95)
        np.testing.assert_allclose(summary.exploratory_p_value, .8025873486341526)

    def test_zero_se_reports_undefined_pvalue(self):
        z, pvalue = I.two_sided_normal_p_value(0, 0)
        self.assertTrue(np.isnan(z) and np.isnan(pvalue))


class ConfidenceIntervalPlotTests(unittest.TestCase):
    def test_intervals_keep_true_bounds_with_points_inside_and_outside(self):
        summary = pd.DataFrame({
            'metric_cn': ['结节级敏感度', 'FP/scan', '阅片时间'],
            'difference': [.1, .3, -.3],
            'ci_lower': [0, .1, -.2], 'ci_upper': [.2, .2, -.1],
            'exploratory_p_value': [.04, .04, .04],
        })
        before = summary.copy(deep=True)
        original_save = Figure.savefig
        saved = []

        def checked_save(figure, path, *args, **kwargs):
            for axis, row in zip(figure.axes, summary.itertuples(index=False)):
                # The interval and point are independent artists: do not clip,
                # widen or reflect the interval to force it around the point.
                interval, point = axis.lines[:2]
                np.testing.assert_allclose(interval.get_xdata(), [row.ci_lower, row.ci_upper])
                np.testing.assert_allclose(interval.get_ydata(), [0, 0])
                np.testing.assert_allclose(point.get_xdata(), [row.difference])
                np.testing.assert_allclose(point.get_ydata(), [0])
            saved.append(Path(path))
            return original_save(figure, path, *args, **kwargs)

        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)/'intervals.png'
            try:
                with patch.object(Figure, 'savefig', autospec=True, side_effect=checked_save):
                    I.plot_confidence_intervals(summary, output)
                self.assertEqual(saved, [output])
                self.assertGreater(output.stat().st_size, 0)
            finally:
                I.plt.close('all')
        pd.testing.assert_frame_equal(summary, before)


class WorkbookExportTests(unittest.TestCase):
    def test_infinities_keep_explicit_marker_and_numeric_cells_stay_numeric(self):
        source = pd.DataFrame({"threshold": [np.inf, .5, -np.inf, np.nan], "TP": [0, 1, 1, 0]})
        before = source.copy(deep=True)
        workbook = Workbook()
        I.add_dataframe_sheet(workbook, "FROC", source, "FROCTest")
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        buffer.seek(0)
        restored = load_workbook(buffer, data_only=True)
        try:
            sheet = restored["FROC"]
            self.assertEqual(sheet["A2"].value, "inf")
            self.assertEqual(sheet["A3"].value, .5)
            self.assertEqual(sheet["A3"].data_type, "n")
            self.assertEqual(sheet["A4"].value, "-inf")
            self.assertIsNone(sheet["A5"].value)
            self.assertEqual(sheet["B3"].value, 1)
        finally:
            restored.close()
        pd.testing.assert_frame_equal(source, before)


if __name__ == "__main__":
    unittest.main()
