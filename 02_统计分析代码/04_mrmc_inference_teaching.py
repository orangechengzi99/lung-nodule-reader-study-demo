from __future__ import annotations

from dataclasses import dataclass
from math import erfc, sqrt
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo
import pandas as pd
from study_config import load_config
from plot_style import configure_plotting, print_plot_style


PATHS, PARAMETERS = load_config(__file__)
ROOT, OUTPUT_DIR, WORKBOOK_PATH = PATHS.root, PATHS.results, PATHS.workbook
DATA_PATH = PATHS.data / "reader_case_results.csv"
BOOTSTRAP_ITERATIONS = PARAMETERS.bootstrap_iterations
RANDOM_SEED = PARAMETERS.bootstrap_seed
CONDITIONS = ("without_ai", "with_ai")

PLOT_STYLE = configure_plotting()


@dataclass(frozen=True)
class MetricResult:
    metric: str
    metric_cn: str
    unit: str
    without_ai: float
    with_ai: float
    difference: float
    bootstrap_se: float
    ci_lower: float
    ci_upper: float
    z_value: float
    exploratory_p_value: float


def validate_input(data: pd.DataFrame) -> tuple[np.ndarray, np.ndarray]:
    required = {
        "reader_id",
        "case_id",
        "condition",
        "period",
        "reference_lesion_count",
        "tp_count",
        "fn_count",
        "fp_count",
        "reading_time_min",
    }
    missing = sorted(required - set(data.columns))
    if missing:
        raise ValueError(f"reader_case_results.csv缺少字段：{missing}")

    if data[list(required)].isna().any().any():
        raise ValueError("推断分析所需字段存在缺失值。")

    count_columns = ["reference_lesion_count", "tp_count", "fn_count", "fp_count"]
    for column in count_columns + ["reading_time_min"]:
        if not data[column].map(pd.api.types.is_number).all():
            raise ValueError(f"{column}必须为数值。")
        values = data[column].to_numpy(dtype=float)
        if not np.isfinite(values).all():
            raise ValueError(f"{column}必须为有限数值。")
        if column in count_columns:
            if not ((values >= 0) & (values == np.floor(values))).all():
                raise ValueError(f"{column}必须为非负整数。")
        elif not (values > 0).all():
            raise ValueError("阅片时间必须大于0。")

    if not (data.groupby("case_id")["reference_lesion_count"].nunique() == 1).all():
        raise ValueError("同一病例的参考结节数必须在所有读者和两种条件下一致。")

    if set(data["condition"].unique()) != set(CONDITIONS):
        raise ValueError("condition必须且只能包含without_ai和with_ai。")

    key = ["reader_id", "case_id", "condition"]
    if data.duplicated(key).any():
        raise ValueError("阅片者－病例－条件组合存在重复记录。")

    if not (
        data["tp_count"] + data["fn_count"]
        == data["reference_lesion_count"]
    ).all():
        raise ValueError("存在TP + FN不等于参考结节数的记录。")

    readers = np.sort(data["reader_id"].unique())
    cases = np.sort(data["case_id"].unique())
    expected_rows = len(readers) * len(cases) * len(CONDITIONS)
    if len(data) != expected_rows:
        raise ValueError(
            "当前教学脚本要求完全交叉数据："
            f"实际{len(data)}行，期望{expected_rows}行。"
        )

    pair_counts = data.groupby(["reader_id", "case_id"])["condition"].nunique()
    if not (pair_counts == len(CONDITIONS)).all():
        raise ValueError("部分阅片者－病例组合没有同时包含两种阅片条件。")

    if not data["period"].isin([1, 2]).all():
        raise ValueError("period必须且只能包含1和2。")
    period_counts = data.groupby(["reader_id", "case_id"])["period"].nunique()
    if not (period_counts == 2).all():
        raise ValueError("部分阅片者－病例组合没有同时包含两个阅片周期。")

    return readers, cases


def build_metric_arrays(
    data: pd.DataFrame,
    readers: np.ndarray,
    cases: np.ndarray,
) -> dict[str, np.ndarray]:
    indexed = data.set_index(["reader_id", "case_id", "condition"]).sort_index()
    arrays: dict[str, np.ndarray] = {}

    for condition in CONDITIONS:
        condition_data = indexed.xs(condition, level="condition")
        condition_data = condition_data.reindex(
            pd.MultiIndex.from_product(
                [readers, cases], names=["reader_id", "case_id"]
            )
        )
        if condition_data.isna().any().any():
            raise ValueError(f"{condition}条件下的数据不能组成完整阅片者×病例矩阵。")

        suffix = "without" if condition == "without_ai" else "with"
        for column in (
            "reference_lesion_count",
            "tp_count",
            "fp_count",
            "reading_time_min",
        ):
            arrays[f"{column}_{suffix}"] = condition_data[column].to_numpy(
                dtype=float
            ).reshape(len(readers), len(cases))

    return arrays


def calculate_reader_averaged_metrics(
    arrays: dict[str, np.ndarray],
    reader_weights: np.ndarray,
    case_weights: np.ndarray,
) -> dict[str, float]:
    expected_shape = arrays["tp_count_without"].shape
    if reader_weights.shape != (expected_shape[0],) or case_weights.shape != (expected_shape[1],):
        raise ValueError("Bootstrap权重长度必须与读者、病例矩阵一致。")
    for weights in (reader_weights, case_weights):
        if not np.isfinite(weights).all() or np.any(weights < 0):
            raise ValueError("Bootstrap权重必须为有限非负数值。")
    reader_weight_total = reader_weights.sum()
    case_weight_total = case_weights.sum()
    if reader_weight_total <= 0 or case_weight_total <= 0:
        raise ValueError("Bootstrap权重总数必须大于0。")

    result: dict[str, float] = {}
    for suffix in ("without", "with"):
        tp_by_reader = arrays[f"tp_count_{suffix}"] @ case_weights
        reference_by_reader = (
            arrays[f"reference_lesion_count_{suffix}"] @ case_weights
        )
        if np.any(reference_by_reader <= 0):
            raise ValueError("重抽病例中没有参考结节，无法计算结节级敏感度。")

        sensitivity_by_reader = tp_by_reader / reference_by_reader
        fp_per_scan_by_reader = (
            arrays[f"fp_count_{suffix}"] @ case_weights / case_weight_total
        )
        reading_time_by_reader = (
            arrays[f"reading_time_min_{suffix}"] @ case_weights
            / case_weight_total
        )

        result[f"lesion_sensitivity_{suffix}"] = float(
            np.average(sensitivity_by_reader, weights=reader_weights)
        )
        result[f"fp_per_scan_{suffix}"] = float(
            np.average(fp_per_scan_by_reader, weights=reader_weights)
        )
        result[f"reading_time_min_{suffix}"] = float(
            np.average(reading_time_by_reader, weights=reader_weights)
        )

    return result


def run_two_way_cluster_bootstrap(
    arrays: dict[str, np.ndarray],
    reader_count: int,
    case_count: int,
) -> tuple[dict[str, float], pd.DataFrame]:
    rng = np.random.default_rng(RANDOM_SEED)
    all_reader_weights = np.ones(reader_count, dtype=float)
    all_case_weights = np.ones(case_count, dtype=float)
    point_estimates = calculate_reader_averaged_metrics(
        arrays,
        all_reader_weights,
        all_case_weights,
    )

    bootstrap_rows: list[dict[str, float | int]] = []
    for iteration in range(1, BOOTSTRAP_ITERATIONS + 1):
        sampled_readers = rng.integers(0, reader_count, size=reader_count)
        sampled_cases = rng.integers(0, case_count, size=case_count)
        reader_weights = np.bincount(
            sampled_readers, minlength=reader_count
        ).astype(float)
        case_weights = np.bincount(sampled_cases, minlength=case_count).astype(float)

        metrics = calculate_reader_averaged_metrics(
            arrays,
            reader_weights,
            case_weights,
        )
        bootstrap_rows.append(
            {
                "iteration": iteration,
                "lesion_sensitivity_difference": (
                    metrics["lesion_sensitivity_with"]
                    - metrics["lesion_sensitivity_without"]
                ),
                "fp_per_scan_difference": (
                    metrics["fp_per_scan_with"]
                    - metrics["fp_per_scan_without"]
                ),
                "reading_time_difference_min": (
                    metrics["reading_time_min_with"]
                    - metrics["reading_time_min_without"]
                ),
            }
        )

    return point_estimates, pd.DataFrame(bootstrap_rows)


def two_sided_normal_p_value(estimate: float, standard_error: float) -> tuple[float, float]:
    if standard_error <= 0:
        return np.nan, np.nan
    z_value = estimate / standard_error
    p_value = erfc(abs(z_value) / sqrt(2.0))
    return float(z_value), float(p_value)


def summarize_inference(
    point_estimates: dict[str, float],
    bootstrap_results: pd.DataFrame,
) -> pd.DataFrame:
    specifications = [
        (
            "lesion_sensitivity",
            "结节级敏感度",
            "proportion",
            "lesion_sensitivity_difference",
        ),
        (
            "fp_per_scan",
            "FP/scan",
            "count_per_scan",
            "fp_per_scan_difference",
        ),
        (
            "reading_time_min",
            "阅片时间",
            "minute",
            "reading_time_difference_min",
        ),
    ]

    rows: list[MetricResult] = []
    for metric, metric_cn, unit, difference_column in specifications:
        without_ai = point_estimates[f"{metric}_without"]
        with_ai = point_estimates[f"{metric}_with"]
        difference = with_ai - without_ai
        distribution = bootstrap_results[difference_column].to_numpy(dtype=float)
        standard_error = float(np.std(distribution, ddof=1))
        ci_lower, ci_upper = np.quantile(distribution, [0.025, 0.975])
        z_value, p_value = two_sided_normal_p_value(difference, standard_error)

        rows.append(
            MetricResult(
                metric=metric,
                metric_cn=metric_cn,
                unit=unit,
                without_ai=without_ai,
                with_ai=with_ai,
                difference=difference,
                bootstrap_se=standard_error,
                ci_lower=float(ci_lower),
                ci_upper=float(ci_upper),
                z_value=z_value,
                exploratory_p_value=p_value,
            )
        )

    summary = pd.DataFrame([row.__dict__ for row in rows])
    summary["ci_contains_zero"] = (
        (summary["ci_lower"] <= 0) & (summary["ci_upper"] >= 0)
    )
    summary["bootstrap_iterations"] = BOOTSTRAP_ITERATIONS
    summary["random_seed"] = RANDOM_SEED
    summary["ci_method"] = "two-way cluster bootstrap percentile interval"
    summary["p_value_method"] = "two-sided normal approximation using bootstrap SE"
    summary["analysis_scope"] = "teaching/exploratory; simulated data"
    return summary


def format_value(metric: str, value: float) -> str:
    if metric == "lesion_sensitivity":
        return f"{value:.1%}"
    if metric == "fp_per_scan":
        return f"{value:.3f}"
    return f"{value:.2f}分钟"


def format_difference(metric: str, value: float) -> str:
    if metric == "lesion_sensitivity":
        return f"{value * 100:+.2f}个百分点"
    if metric == "fp_per_scan":
        return f"{value:+.3f}"
    return f"{value:+.2f}分钟"


def format_interval(metric: str, lower: float, upper: float) -> str:
    if metric == "lesion_sensitivity":
        return f"{lower * 100:.2f}～{upper * 100:.2f}个百分点"
    if metric == "fp_per_scan":
        return f"{lower:.3f}～{upper:.3f}"
    return f"{lower:.2f}～{upper:.2f}分钟"


def write_markdown_summary(summary: pd.DataFrame, path: Path) -> None:
    lines = [
        "# MRMC教学性推断结果说明",
        "",
        "## 方法",
        "",
        (
            f"本分析使用{BOOTSTRAP_ITERATIONS}次阅片者－病例双重Bootstrap。"
            "每次分别对阅片者和病例进行有放回抽样，并保留同一阅片者、同一病例"
            "在无AI与AI辅助条件下的配对关系。"
        ),
        "",
        (
            "95%置信区间采用Bootstrap百分位法；探索性P值采用点估计除以"
            "Bootstrap标准误形成的双侧正态近似检验。"
        ),
        "",
        "## 结果",
        "",
        "| 指标 | 无AI | AI辅助 | 差异（AI辅助－无AI） | 95% CI | 探索性P值 |",
        "|---|---:|---:|---:|---:|---:|",
    ]

    for row in summary.itertuples(index=False):
        p_text = "<0.0001" if row.exploratory_p_value < 0.0001 else f"{row.exploratory_p_value:.4f}"
        lines.append(
            "| "
            f"{row.metric_cn} | "
            f"{format_value(row.metric, row.without_ai)} | "
            f"{format_value(row.metric, row.with_ai)} | "
            f"{format_difference(row.metric, row.difference)} | "
            f"{format_interval(row.metric, row.ci_lower, row.ci_upper)} | "
            f"{p_text} |"
        )

    lines.extend(
        [
            "",
            "## 解释边界",
            "",
            "- 数据完全由程序生成，不来源于真实患者、真实阅片者或真实AI产品。",
            "- 本分析用于演示阅片者和病例抽样变异如何影响标准误、置信区间和P值。",
            "- 当前方法属于教学性、探索性分析，不替代预先规定并经统计师确认的正式MRMC主分析。",
            "- 当前数据仅保留每个阅片者－病例－条件下的结节计数，没有逐结节坐标和匹配记录，不能作为完整定位型JAFROC分析。",
            "- 即使探索性P值较小，也不能据此宣称真实AI产品具有临床有效性。",
            "",
            "## 可复核信息",
            "",
            f"- 随机种子：`{RANDOM_SEED}`",
            f"- Bootstrap次数：`{BOOTSTRAP_ITERATIONS}`",
            f"- 输入：`{DATA_PATH.relative_to(ROOT).as_posix()}`",
            (
                "- 明细：`"
                f"{(OUTPUT_DIR / 'mrmc_bootstrap_distributions.csv').relative_to(ROOT).as_posix()}`"
            ),
            (
                "- 汇总：`"
                f"{(OUTPUT_DIR / 'mrmc_inference_summary.csv').relative_to(ROOT).as_posix()}`"
            ),
        ]
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def add_dataframe_sheet(
    workbook: Workbook,
    title: str,
    data: pd.DataFrame,
    table_name: str,
) -> None:
    sheet = workbook.create_sheet(title)
    clean_data = data.astype(object).where(pd.notna(data), None)
    for row in dataframe_to_rows(clean_data, index=False, header=True):
        # XLSX cannot store numeric infinity: preserve the sentinel explicitly.
        # Finite values remain numeric; missing values remain genuinely blank.
        excel_row = [
            ("inf" if value > 0 else "-inf")
            if isinstance(value, (float, np.floating)) and np.isinf(value)
            else value
            for value in row
        ]
        sheet.append(excel_row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    table = Table(displayName=table_name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)

    for column_index, column_cells in enumerate(sheet.columns, start=1):
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        width = min(max([len(value) for value in values] + [10]) + 2, 36)
        sheet.column_dimensions[get_column_letter(column_index)].width = width
        for cell in column_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def build_results_workbook(
    summary: pd.DataFrame,
    bootstrap_results: pd.DataFrame,
) -> None:
    data_dir = DATA_PATH.parent
    required_tables = [
        ("OverallSummary", OUTPUT_DIR / "overall_descriptive_summary.csv"),
        ("FROCResults", OUTPUT_DIR / "froc_results.csv"),
        ("PairedDifferences", OUTPUT_DIR / "paired_differences.csv"),
        ("ReaderSummary", OUTPUT_DIR / "reader_condition_summary.csv"),
        ("QualityChecks", OUTPUT_DIR / "quality_checks.csv"),
        ("Readers", data_dir / "readers.csv"),
        ("Cases", data_dir / "cases.csv"),
        ("ReaderCaseResults", data_dir / "reader_case_results.csv"),
        ("FROCScans", data_dir / "froc_scans.csv"),
        ("FROCReferences", data_dir / "froc_references.csv"),
        ("FROCCandidates", data_dir / "froc_candidates.csv"),
        ("DataDictionary", data_dir / "data_dictionary.csv"),
    ]
    missing = [str(path.relative_to(ROOT)) for _, path in required_tables if not path.exists()]
    if missing:
        raise FileNotFoundError(
            "生成综合工作簿前缺少以下文件，请按README顺序运行脚本："
            f"{missing}"
        )

    workbook = Workbook()
    readme = workbook.active
    readme.title = "ReadMe"
    readme.merge_cells("A1:H1")
    readme["A1"] = "MRMC辅助阅片教学数据与结果"
    readme["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    readme["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    readme["A1"].alignment = Alignment(horizontal="center", vertical="center")
    readme.row_dimensions[1].height = 28

    readme_rows = [
        ("数据性质", "完全虚构的教学数据，不来源于真实患者、医院、阅片者或AI产品。"),
        ("研究结构", "8名阅片者、80例CT、无AI和AI辅助两种条件，共1280条判读记录。"),
        ("主要指标", "阅片者平均结节级敏感度差异。"),
        ("关键次要终点", "阅片者平均FP/scan差异。"),
        ("统计范围", "描述性汇总、FROC教学分析及阅片者－病例双重Bootstrap教学性推断。"),
        ("结论边界", "不能作为真实产品临床有效性、临床获益或注册证据。"),
        ("更新方式", "依次运行01至04脚本；04脚本在完成推断后刷新本工作簿。"),
    ]
    readme.append([])
    readme.append(["项目", "说明"])
    for item, description in readme_rows:
        readme.append([item, description])
    for cell in readme[3]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    readme.column_dimensions["A"].width = 20
    readme.column_dimensions["B"].width = 88
    for row in readme.iter_rows(min_row=4, max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    add_dataframe_sheet(workbook, "MRMCInference", summary, "T01MRMCInference")
    add_dataframe_sheet(
        workbook,
        "BootstrapDistributions",
        bootstrap_results,
        "T02Bootstrap",
    )
    for index, (title, path) in enumerate(required_tables, start=3):
        add_dataframe_sheet(
            workbook,
            title,
            pd.read_csv(path),
            f"T{index:02d}{title}",
        )

    WORKBOOK_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(WORKBOOK_PATH)


def plot_confidence_intervals(summary: pd.DataFrame, path: Path) -> None:
    fig, axes = plt.subplots(3, 1, figsize=(9, 9))
    fig.suptitle(PLOT_STYLE.text("MRMC教学性推断：AI辅助减无AI"), fontsize=16, fontweight="bold")

    colors = ["#0072B2", "#D55E00", "#009E73"]
    x_labels = ["敏感度差异", "FP/scan差异", "阅片时间差异（分钟）"]
    for axis, row, color, x_label in zip(
        axes,
        summary.itertuples(index=False),
        colors,
        x_labels,
    ):
        # A percentile interval need not contain the point estimate. Draw its
        # true endpoints independently, without changing any statistical values.
        axis.plot(
            [row.ci_lower, row.ci_upper],
            [0, 0],
            marker="|",
            markersize=12,
            color=color,
            linewidth=2,
        )
        axis.plot(
            row.difference,
            0,
            marker="o",
            color=color,
            markersize=8,
            linestyle="none",
        )
        axis.axvline(0, color="#555555", linestyle="--", linewidth=1.2)
        axis.set_yticks([])
        axis.set_xlabel(PLOT_STYLE.text(x_label))
        axis.grid(axis="x", alpha=0.25)
        p_text = (
            "<0.0001"
            if row.exploratory_p_value < 0.0001
            else f"={row.exploratory_p_value:.4f}"
        )
        if PLOT_STYLE.language == "zh":
            title = (f"{row.metric_cn}：差异={row.difference:.4f}，"
                     f"95% CI {row.ci_lower:.4f}～{row.ci_upper:.4f}，探索性P{p_text}")
        else:
            title = (f"{PLOT_STYLE.text(row.metric_cn)}: difference={row.difference:.4f}, "
                     f"95% CI {row.ci_lower:.4f} to {row.ci_upper:.4f}, exploratory P{p_text}")
        axis.set_title(title, fontsize=11)

    fig.text(
        0.5,
        0.01,
        PLOT_STYLE.text("虚构模拟数据；双重Bootstrap教学演示，不作为正式临床证据"),
        ha="center",
        color="#555555",
        fontsize=10,
    )
    fig.tight_layout(rect=(0, 0.04, 1, 0.96))
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)


def plot_workbook_summary_preview(path: Path) -> None:
    summary_path = OUTPUT_DIR / "overall_descriptive_summary.csv"
    if not summary_path.exists():
        raise FileNotFoundError(
            "生成工作簿预览图前缺少overall_descriptive_summary.csv，请先运行02脚本。"
        )

    summary = pd.read_csv(summary_path)
    display = summary.rename(
        columns={
            "without_ai_reader_mean": "without_ai",
            "with_ai_reader_mean": "with_ai",
            "mean_paired_difference": "paired_difference",
            "minimum_reader_difference": "min_difference",
            "maximum_reader_difference": "max_difference",
        }
    ).copy()
    numeric_columns = display.select_dtypes(include="number").columns
    display[numeric_columns] = display[numeric_columns].round(3)

    fig, axis = plt.subplots(figsize=(17, 2.8))
    axis.axis("off")
    table = axis.table(
        cellText=display.astype(str).values,
        colLabels=display.columns,
        colWidths=[0.10, 0.20, 0.10, 0.10, 0.11, 0.11, 0.11, 0.17],
        cellLoc="center",
        loc="center",
    )
    table.auto_set_font_size(False)
    table.set_fontsize(7)
    table.scale(1, 1.45)

    for (row, _), cell in table.get_celld().items():
        if row == 0:
            cell.set_facecolor("#1F4E78")
            cell.set_text_props(color="white", fontweight="bold")
        elif row % 2 == 0:
            cell.set_facecolor("#D9EEF7")

    fig.tight_layout(pad=0.4)
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)


def main() -> None:
    print_plot_style(PLOT_STYLE)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    data = pd.read_csv(DATA_PATH)
    readers, cases = validate_input(data)
    arrays = build_metric_arrays(data, readers, cases)

    point_estimates, bootstrap_results = run_two_way_cluster_bootstrap(
        arrays,
        reader_count=len(readers),
        case_count=len(cases),
    )
    summary = summarize_inference(point_estimates, bootstrap_results)

    summary.to_csv(
        OUTPUT_DIR / "mrmc_inference_summary.csv",
        index=False,
        encoding="utf-8-sig",
    )
    bootstrap_results.to_csv(
        OUTPUT_DIR / "mrmc_bootstrap_distributions.csv",
        index=False,
        encoding="utf-8-sig",
    )
    write_markdown_summary(
        summary,
        OUTPUT_DIR / "MRMC教学性推断结果说明.md",
    )
    plot_confidence_intervals(
        summary,
        OUTPUT_DIR / "mrmc_inference_confidence_intervals.png",
    )
    build_results_workbook(summary, bootstrap_results)
    plot_workbook_summary_preview(
        OUTPUT_DIR / "workbook_summary_preview.png",
    )

    display_columns = [
        "metric_cn",
        "without_ai",
        "with_ai",
        "difference",
        "bootstrap_se",
        "ci_lower",
        "ci_upper",
        "exploratory_p_value",
    ]
    print("MRMC教学性推断结果（虚构模拟数据）")
    print(summary[display_columns].to_string(index=False))
    print(f"\n阅片者数：{len(readers)}")
    print(f"病例数：{len(cases)}")
    print(f"Bootstrap次数：{BOOTSTRAP_ITERATIONS}")
    print(f"结果目录：{OUTPUT_DIR}")
    print(f"综合工作簿：{WORKBOOK_PATH}")


if __name__ == "__main__":
    main()
